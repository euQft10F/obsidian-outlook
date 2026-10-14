#!/usr/bin/env python3
"""Convert Outlook .msg (and .eml) files from 00-inbox/ into immutable markdown notes in raw/.

Deterministic, no LLM. Safe to run repeatedly: duplicates (same Message-ID or same bytes)
are detected and skipped.

    python msg_to_md.py                 # convert everything in <vault>/00-inbox
    python msg_to_md.py --dry-run       # show what would happen
    python msg_to_md.py --keep          # leave originals in the inbox (default: archive to raw/msg/)
    python msg_to_md.py --me me@work.com --me me@home.com   # mark emails you sent

Output per email:
    raw/emails/YYYY/MM/<YYYY-MM-DD HHMM> <Sender> - <Subject>.md
    raw/attachments/YYYY/MM/<same stem>/<files>   (+ <file>.extracted.md text for pdf/docx/xlsx/txt/csv)
    raw/msg/YYYY/MM/<same stem>.msg               (archived original)

Requires: extract-msg (for .msg). Optional: markdownify (HTML-only emails), pypdf,
python-docx, openpyxl (attachment text extraction).
"""
from __future__ import annotations

import argparse
import datetime as dt
import email
import email.policy
import email.utils
import hashlib
import json
import re
import shutil
import sys
import traceback
from dataclasses import dataclass, field
from enum import StrEnum
from html.parser import HTMLParser
from pathlib import Path

from vaultlib import (
    EMAIL_SUFFIXES,
    FAILED_DIRNAME,
    INBOX_DIRNAME,
    check_setup,
    find_vault,
    owner_addresses,
    sanitize_filename,
    utf8_stdout,
)

INDEX_FILE = "raw/.email-index.json"
LOG_FILE = "raw/.ingest-log.jsonl"

TEXT_EXTRACT_LIMIT = 200_000  # chars kept from an attachment's extracted text
MAX_EXTRACT_BYTES = 25 * 1024 * 1024  # attachments above this are saved but never parsed
SMALL_INLINE_IMAGE = 20_000  # bytes; smaller inline images are usually signature logos
SHORT_BODY_CHARS = 20  # a plain-text body shorter than this is replaced by the HTML body
XLSX_MAX_ROWS = 500
MAX_ATTACHMENT_NAME = 120
CONVERSATION_ROOT_BYTES = 22  # the first 22 bytes of an Outlook ConversationIndex identify the thread
MAX_UNIQUE_SUFFIX = 1000

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".emz", ".wmz", ".tif", ".tiff", ".webp"}
PLAIN_TEXT_EXT = {".txt", ".csv", ".md", ".json", ".log"}
# Executable / script attachments are saved with a ".blocked" suffix and never parsed.
BLOCKED_EXT = {".exe", ".dll", ".js", ".jse", ".vbs", ".vbe", ".ps1", ".bat", ".cmd", ".lnk", ".scr",
               ".jar", ".msi", ".hta", ".com", ".cpl", ".wsf", ".reg"}
# Attachments Obsidian would treat as vault notes or views get ".txt" appended, so an emailed
# `Priya Shah.md` can't pose as a wiki page or capture [[links]]; the text is still extracted.
NOTE_LIKE_EXT = {".md", ".canvas", ".base"}

SUBJECT_PREFIX_RE = re.compile(
    r"^\s*((re|fw|fwd|aw|wg|tr|sv|vs|rv|r|antw|odp|réf|ref)\s*(\[\d+\])?\s*:\s*|\[external\]\s*)+",
    re.I,
)
BARE_ADDRESS_RE = re.compile(r"[\w.+'&-]+@[\w-]+(?:\.[\w-]+)+")
MESSAGE_ID_RE = re.compile(r"<[^<>\s]+>")

# Where quoted history starts (Outlook, Gmail, Apple Mail; EN/FR/DE/NL/IT/ES headers).
QUOTE_START_PATTERNS = [
    re.compile(r"^-{2,}\s*(original message|ursprüngliche nachricht|message d'origine|oorspronkelijk bericht|messaggio originale|mensaje original)\s*-{2,}\s*$", re.I),  # noqa: E501
    re.compile(r"^_{10,}\s*$"),
    re.compile(r"^(on|le|am|op|il|el)\s.{4,250}(wrote|a écrit|schrieb|schreef|ha scritto|escribió)\s*:?\s*$", re.I),
]
HEADER_FROM_RE = re.compile(r"^\**\s*(from|de|von|van|da|od)\s*:\**\s*\S", re.I)
HEADER_SENT_RE = re.compile(r"^\**\s*(sent|date|envoyé|gesendet|verzonden|inviato|enviado|datum)\s*:\**\s*\S", re.I)
HEADER_LOOKAHEAD_LINES = 5  # a "From:" line starts a quoted header block if "Sent:" follows this closely
QUOTE_RUN_MIN_LINES = 3  # a run of ">" lines this long (and mostly quoted to the end) is reply quoting
QUOTE_RUN_MIN_RATIO = 0.6


class DateSource(StrEnum):
    HEADER = "header"
    FILE_MODIFIED_TIME = "file-modified-time"
    PARENT_EMAIL = "parent-email"


class RefusedFilesError(Exception):
    """Explicitly named files that are not inside the vault's 00-inbox/."""


# ---------------------------------------------------------------- text helpers

def normalize_text(value) -> str:
    """Any value -> str with NULs removed and newlines normalised, stripped."""
    if value is None:
        return ""
    if isinstance(value, bytes):
        value = value.decode("utf-8", errors="replace")
    text = str(value).replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    return text.strip()


def optional_attr(obj, name: str, default=None):
    """`getattr` that also swallows errors from lazily-parsed library properties (extract_msg raises freely)."""
    try:
        value = getattr(obj, name)
    except Exception:
        return default
    return default if value is None else value


def normalize_subject(subject: str) -> str:
    """Subject without RE:/FW:/[External] prefixes; the key used to group threads."""
    stripped = SUBJECT_PREFIX_RE.sub("", subject or "").strip()
    return stripped or "(no subject)"


def slugify(text: str, max_length: int = 60) -> str:
    slug = re.sub(r"[^\w]+", "-", text.lower(), flags=re.U).strip("-")
    return slug[:max_length].strip("-") or "untitled"


def one_line(text: str) -> str:
    """Collapse all whitespace (including newlines) so header text can't start new markdown lines."""
    return " ".join(text.split())


def defang_markdown(text: str) -> str:
    """Email text as Obsidian will render it: no raw HTML and no `![](…)` embeds, which load remote
    content (tracking pixels, iframes) when the note is opened. Links stay as plain, click-only text."""
    return text.replace("<", "&lt;").replace("![", "!\\[")


def yaml_str(value) -> str:
    return json.dumps("" if value is None else str(value), ensure_ascii=False)


def yaml_list(items) -> str:
    items = [item for item in items if item]
    if not items:
        return " []"
    return "\n" + "\n".join(f"  - {yaml_str(item)}" for item in items)


def message_ids(text: str) -> list[str]:
    return MESSAGE_ID_RE.findall(normalize_text(text))


def first_message_id(text: str) -> str:
    return (message_ids(text) or [""])[0]


# ---------------------------------------------------------------- addresses

@dataclass(frozen=True, slots=True)
class Address:
    name: str
    email: str

    def display(self) -> str:
        if self.email and self.name and self.name.lower() != self.email:
            return f"{self.name} <{self.email}>"
        return self.email or self.name


UNKNOWN_SENDER = Address("Unknown sender", "")


def _peel_leading_addresses(name: str) -> tuple[list[str], str]:
    """Bare addresses that precede a display name: "a@b.com, Smith, John" -> (["a@b.com"], "Smith, John")."""
    leading = []
    while match := re.match(r"^\s*(" + BARE_ADDRESS_RE.pattern + r")\s*,\s*", name):
        leading.append(match.group(1))
        name = name[match.end():]
    return leading, name


def _split_recipients(raw: str) -> list[tuple[str, str]]:
    """Raw (name, address) pairs, before cleaning; either part may be empty."""
    pairs: list[tuple[str, str]] = []
    for segment in re.split(r"(?<=>)\s*[,;]\s*|\s*;\s*", raw):
        segment = segment.strip(" ,")
        if not segment:
            continue
        bracketed = re.search(r"<([^<>]+)>\s*$", segment)
        if bracketed:
            leading, name = _peel_leading_addresses(segment[:bracketed.start()].strip())
            pairs.extend(("", address) for address in leading)
            pairs.append((name, bracketed.group(1)))
        elif found := BARE_ADDRESS_RE.findall(segment):
            pairs.extend(("", address) for address in found)
        else:
            pairs.append((segment, ""))  # Exchange display name with no address
    return pairs


def _clean_pair(name: str, address: str) -> tuple[str, str]:
    name = normalize_text(name).strip("'\" ,")
    address = normalize_text(address).strip().lower().removeprefix("mailto:")
    if address and "@" not in address:  # Exchange X.500 path (/o=exchangelabs/...): keep the name only
        address = ""
    if "@" in name and not address:
        name, address = "", name.lower()
    if name.lower() == address:
        name = ""
    return name, address


def parse_addresses(raw: str) -> list[Address]:
    """Parse Outlook-style recipient strings.

    Handles 'Last, First <a@b>; Other <c@d>', unquoted names containing commas or '@',
    and bare address lists. (email.utils.getaddresses mangles the Outlook forms.)
    """
    addresses: list[Address] = []
    seen: set[str] = set()
    for raw_name, raw_address in _split_recipients(normalize_text(raw)):
        name, address = _clean_pair(raw_name, raw_address)
        key = address or name.lower()
        if not key or key in seen:
            continue
        seen.add(key)
        addresses.append(Address(name or (address.split("@")[0] if address else ""), address))
    return addresses


# ---------------------------------------------------------------- body text

class _TextExtractor(HTMLParser):
    """Fallback HTML -> text when markdownify is not installed."""

    BLOCK_TAGS = ("br", "p", "div", "tr", "li", "h1", "h2", "h3", "table")
    SKIPPED_TAGS = ("script", "style", "head")

    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIPPED_TAGS:
            self._skip_depth += 1
        if tag in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIPPED_TAGS and self._skip_depth:
            self._skip_depth -= 1

    def handle_data(self, data):
        if not self._skip_depth:
            self.parts.append(data)


def html_to_markdown(html_text: str) -> str:
    try:
        from markdownify import markdownify  # type: ignore
        markdown = markdownify(html_text, heading_style="ATX", strip=["img", "style", "script"])
    except Exception:  # markdownify missing or choking on odd HTML: fall back to plain text
        extractor = _TextExtractor()
        extractor.feed(html_text)
        markdown = "".join(extractor.parts)  # entities already decoded; decoding twice revives `&lt;tags&gt;`
    markdown = re.sub(r"[ \t]+\n", "\n", markdown)
    return re.sub(r"\n{3,}", "\n\n", markdown).strip()


def _quoted_history_starts_at(lines: list[str], index: int, quoted_lines: int, nonblank_lines: int) -> bool:
    """`quoted_lines` / `nonblank_lines` count the ">"-lines and non-blank lines from `index` to the end."""
    stripped = lines[index].strip()
    if any(pattern.match(stripped) for pattern in QUOTE_START_PATTERNS):
        return True
    if HEADER_FROM_RE.match(stripped):
        following = lines[index + 1:index + 1 + HEADER_LOOKAHEAD_LINES]
        if any(HEADER_SENT_RE.match(line.strip()) for line in following):
            return True
    if stripped.startswith(">"):
        return quoted_lines >= QUOTE_RUN_MIN_LINES and quoted_lines >= QUOTE_RUN_MIN_RATIO * nonblank_lines
    return False


def split_quoted(body: str) -> tuple[str, str]:
    """Split body into (new content, quoted history)."""
    lines = body.split("\n")
    nonblank_from = [0] * (len(lines) + 1)  # suffix counts, so each ">" line is judged in O(1)
    quoted_from = [0] * (len(lines) + 1)
    for index in range(len(lines) - 1, -1, -1):
        stripped = lines[index].strip()
        nonblank_from[index] = nonblank_from[index + 1] + bool(stripped)
        quoted_from[index] = quoted_from[index + 1] + stripped.startswith(">")
    for index, line in enumerate(lines):
        if index == 0 or not line.strip():
            continue
        if _quoted_history_starts_at(lines, index, quoted_from[index], nonblank_from[index]):
            return "\n".join(lines[:index]).rstrip(), "\n".join(lines[index:]).strip()
    return body.strip(), ""


# ---------------------------------------------------------------- attachments

def _extract_plain(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader  # type: ignore
    return "\n\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)


def _extract_docx(path: Path) -> str:
    import docx  # type: ignore
    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    import openpyxl  # type: ignore
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"## Sheet: {sheet.title}")
        for row_number, row in enumerate(sheet.iter_rows(values_only=True)):
            if row_number >= XLSX_MAX_ROWS:
                lines.append("… (truncated)")
                break
            if any(cell is not None for cell in row):
                lines.append(" | ".join("" if cell is None else str(cell) for cell in row))
    return "\n".join(lines)


TEXT_EXTRACTORS = {
    **dict.fromkeys(PLAIN_TEXT_EXT, _extract_plain),
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    **dict.fromkeys((".xlsx", ".xlsm"), _extract_xlsx),
}


def extract_attachment_text(path: Path) -> str:
    """Text of a pdf/docx/xlsx/plain attachment; "" if unsupported or the parser library is missing."""
    extractor = TEXT_EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        return ""
    try:
        return extractor(path)
    except ImportError:
        return ""
    except Exception as error:  # unreadable / encrypted attachment
        return f"(text extraction failed: {error})"


# ---------------------------------------------------------------- parsing

@dataclass(slots=True)
class Attachment:
    filename: str
    data: bytes
    inline: bool


@dataclass(slots=True)
class ParsedEmail:
    """Normalised view of one email, independent of source format."""
    subject: str = ""
    sender: list[Address] = field(default_factory=list)
    to: list[Address] = field(default_factory=list)
    cc: list[Address] = field(default_factory=list)
    date: dt.datetime | None = None
    date_source: DateSource = DateSource.HEADER
    message_id: str = ""
    in_reply_to: str = ""
    references: list[str] = field(default_factory=list)
    conversation_root: str = ""
    body: str = ""
    attachments: list[Attachment] = field(default_factory=list)
    embedded: list[ParsedEmail] = field(default_factory=list)
    importance: str = ""
    categories: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _msg_threading_headers(parsed: ParsedEmail, message) -> None:
    """References / In-Reply-To / Date from the raw transport header, for fields extract_msg left empty."""
    header = optional_attr(message, "header")
    if header is None:
        return
    try:
        parsed.references = message_ids(header.get("References", "") or "")
        if not parsed.in_reply_to:
            parsed.in_reply_to = first_message_id(header.get("In-Reply-To", "") or "")
        if not parsed.date and header.get("Date"):
            parsed.date = email.utils.parsedate_to_datetime(header.get("Date"))
    except (AttributeError, TypeError, ValueError):
        pass  # malformed transport header: keep whatever was already found


def _msg_attachments(parsed: ParsedEmail, message) -> None:
    for number, attachment in enumerate(optional_attr(message, "attachments", []) or [], start=1):
        try:
            data = attachment.data
            if not isinstance(data, (bytes, bytearray, str)) and hasattr(data, "attachments") and hasattr(data, "body"):
                parsed.embedded.append(parse_msg_object(data))
                continue
            if not isinstance(data, (bytes, bytearray)) or not data:
                continue
            filename = normalize_text(optional_attr(attachment, "longFilename")
                                      or optional_attr(attachment, "shortFilename")
                                      or f"attachment-{number}")
            inline = bool(optional_attr(attachment, "hidden", False)) \
                or bool(normalize_text(optional_attr(attachment, "contentId", "")))
            parsed.attachments.append(Attachment(filename, bytes(data), inline))
        except Exception as error:  # a damaged attachment must not lose the whole email
            parsed.warnings.append(f"attachment {number} skipped: {type(error).__name__}: {error}")


def parse_msg_object(message) -> ParsedEmail:
    parsed = ParsedEmail(subject=normalize_text(optional_attr(message, "subject", "")))
    parsed.sender = parse_addresses(optional_attr(message, "sender", ""))
    parsed.to = parse_addresses(optional_attr(message, "to", ""))
    parsed.cc = parse_addresses(optional_attr(message, "cc", ""))
    date = optional_attr(message, "date")
    parsed.date = date if isinstance(date, dt.datetime) else None
    parsed.message_id = first_message_id(optional_attr(message, "messageId", ""))
    parsed.in_reply_to = first_message_id(optional_attr(message, "inReplyTo", ""))
    _msg_threading_headers(parsed, message)

    conversation_index = optional_attr(message, "conversationIndex")
    if isinstance(conversation_index, (bytes, bytearray)) and len(conversation_index) >= CONVERSATION_ROOT_BYTES:
        parsed.conversation_root = conversation_index[:CONVERSATION_ROOT_BYTES].hex()
    importance = optional_attr(message, "importance")
    if importance is not None and "HIGH" in str(importance).upper():
        parsed.importance = "high"
    categories = optional_attr(message, "categories")
    if categories:
        names = categories if isinstance(categories, (list, tuple)) else [categories]
        parsed.categories = [normalize_text(name) for name in names if normalize_text(name)]

    body = normalize_text(optional_attr(message, "body", ""))
    if len(body) < SHORT_BODY_CHARS:
        html_body = optional_attr(message, "htmlBody")
        if html_body:
            body = html_to_markdown(normalize_text(html_body)) or body
    parsed.body = body
    _msg_attachments(parsed, message)
    return parsed


def parse_msg_file(path: Path) -> ParsedEmail:
    import extract_msg  # type: ignore

    message = extract_msg.openMsg(str(path))
    try:
        return parse_msg_object(message)
    finally:
        try:
            message.close()
        except Exception:  # best-effort cleanup of the OLE handle
            pass


def _parse_eml_message(message) -> ParsedEmail:
    parsed = ParsedEmail(subject=normalize_text(message.get("Subject", "")))
    parsed.sender = parse_addresses(str(message.get("From", "")))
    parsed.to = parse_addresses(str(message.get("To", "")))
    parsed.cc = parse_addresses(str(message.get("Cc", "")))
    try:
        parsed.date = email.utils.parsedate_to_datetime(str(message.get("Date"))) if message.get("Date") else None
    except (TypeError, ValueError):
        parsed.date = None
    parsed.message_id = first_message_id(str(message.get("Message-ID", "")))
    parsed.in_reply_to = first_message_id(str(message.get("In-Reply-To", "")))
    parsed.references = message_ids(str(message.get("References", "")))
    if "high" in str(message.get("Importance", "")).lower() or str(message.get("X-Priority", "")).startswith("1"):
        parsed.importance = "high"

    plain = message.get_body(preferencelist=("plain",))
    body = normalize_text(plain.get_content()) if plain else ""
    if len(body) < SHORT_BODY_CHARS:
        html_part = message.get_body(preferencelist=("html",))
        if html_part:
            body = html_to_markdown(html_part.get_content()) or body
    parsed.body = body

    for number, part in enumerate(message.iter_attachments(), start=1):
        try:
            if part.get_content_type() == "message/rfc822":
                inner = part.get_payload(0) if part.is_multipart() else None
                if inner is not None:
                    reparsed = email.message_from_bytes(inner.as_bytes(), policy=email.policy.default)
                    parsed.embedded.append(_parse_eml_message(reparsed))
                continue
            data = part.get_payload(decode=True)
            if not data:
                continue
            filename = normalize_text(part.get_filename() or f"attachment-{number}")
            inline = part.get_content_disposition() == "inline" or bool(part.get("Content-ID"))
            parsed.attachments.append(Attachment(filename, data, inline))
        except Exception as error:  # a damaged part must not lose the whole email
            parsed.warnings.append(f"attachment {number} skipped: {type(error).__name__}: {error}")
    return parsed


def parse_eml_file(path: Path) -> ParsedEmail:
    return _parse_eml_message(email.message_from_bytes(path.read_bytes(), policy=email.policy.default))


# ---------------------------------------------------------------- rendering

def thread_id_for(parsed: ParsedEmail) -> str:
    """Stable thread key: Outlook conversation index, else the thread root's Message-ID, else the subject."""
    if parsed.conversation_root:
        return "conv-" + parsed.conversation_root[:24]
    # References[0] is the thread root; a root email's own Message-ID matches its replies' References[0].
    root = parsed.references[0] if parsed.references else (parsed.in_reply_to or parsed.message_id)
    if root:
        return "ref-" + hashlib.sha1(root.encode()).hexdigest()[:16]
    return "subj-" + slugify(normalize_subject(parsed.subject), 48)


def render_note(parsed: ParsedEmail, *, sender: Address, local: dt.datetime, direction: str,
                attachment_lines: list[str], source_format: str, source_file: str, sha: str,
                embedded_in: str) -> str:
    """The complete markdown note: frontmatter, headers, new text, attachments, folded quoted history."""
    subject = parsed.subject or "(no subject)"
    participants = sorted({address.email for address in parsed.sender + parsed.to + parsed.cc if address.email})
    frontmatter = [
        "---",
        "type: email",
        f"subject: {yaml_str(subject)}",
        f"thread_subject: {yaml_str(normalize_subject(parsed.subject))}",
        f"thread_id: {yaml_str(thread_id_for(parsed))}",
        f"date: {local:%Y-%m-%dT%H:%M:%S}",
        f"date_source: {parsed.date_source}",
        f"from: {yaml_str(sender.display())}",
        f"from_email: {yaml_str(sender.email)}",
        f"to:{yaml_list(address.display() for address in parsed.to)}",
        f"cc:{yaml_list(address.display() for address in parsed.cc)}",
        f"participants:{yaml_list(participants)}",
        f"direction: {direction}",
        f"has_attachments: {'true' if attachment_lines else 'false'}",
    ]
    if parsed.importance:
        frontmatter.append(f"importance: {parsed.importance}")
    if parsed.categories:
        frontmatter.append(f"outlook_categories:{yaml_list(parsed.categories)}")
    if embedded_in:
        frontmatter.append(f"embedded_in: {yaml_str('[[' + embedded_in + ']]')}")
    frontmatter += [
        f"message_id: {yaml_str(parsed.message_id)}",
        f"in_reply_to: {yaml_str(parsed.in_reply_to)}",
        f"source_format: {source_format}",
        f"source_file: {yaml_str(source_file)}",
        f"sha256: {sha[:16]}",
        "wiki_status: pending",
        "tags:\n  - email",
        "---",
    ]

    new_text, quoted = (defang_markdown(part) for part in split_quoted(parsed.body))

    def header(addresses: list[Address]) -> str:
        return defang_markdown(one_line(", ".join(address.display() for address in addresses)))

    lines = frontmatter + [
        "",
        f"# {defang_markdown(one_line(subject))}",
        "",
        f"**From:** {header([sender])}  ",
        f"**To:** {header(parsed.to) or '—'}  ",
    ]
    if parsed.cc:
        lines.append(f"**Cc:** {header(parsed.cc)}  ")
    lines += [f"**Date:** {local:%a %d %b %Y %H:%M}", "", "## Message", "", new_text or "_(empty body)_", ""]
    if attachment_lines:
        lines += ["## Attachments", "", *attachment_lines, ""]
    if quoted:
        lines += ["## Quoted history", "", "> [!quote]- Earlier messages quoted in this email (folded)"]
        lines += [("> " + line) if line.strip() else ">" for line in quoted.split("\n")]
        lines.append("")
    return "\n".join(lines)


# ---------------------------------------------------------------- converting

@dataclass(slots=True)
class ConversionResult:
    file: str
    status: str  # converted | duplicate | failed
    note: str | None = None
    error: str | None = None
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        data = {"file": self.file, "status": self.status}
        if self.note is not None:
            data["note"] = self.note
        if self.error is not None:
            data["error"] = self.error
        if self.warnings:
            data["warnings"] = self.warnings
        return data


class ConversionIndex:
    """What has already been converted: note path by Message-ID and by file hash (raw/.email-index.json)."""

    def __init__(self, path: Path):
        self.path = path
        self.data: dict[str, dict[str, str]] = {"by_msgid": {}, "by_sha": {}}
        if path.exists():
            self.data.update(json.loads(path.read_text(encoding="utf-8")))

    def note_for_sha(self, sha: str) -> str | None:
        return self.data["by_sha"].get(sha)

    def note_for_message_id(self, message_id: str) -> str | None:
        return self.data["by_msgid"].get(message_id) if message_id else None

    def add(self, note: str, *, message_id: str, sha: str | None) -> None:
        if message_id:
            self.data["by_msgid"][message_id] = note
        if sha:
            self.data["by_sha"][sha] = note

    def save(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        temporary = self.path.with_name(self.path.name + ".tmp")
        temporary.write_text(json.dumps(self.data, indent=1, ensure_ascii=False), encoding="utf-8")
        temporary.replace(self.path)


class Converter:
    def __init__(self, vault: Path, own_addresses: set[str], *, dry_run: bool, keep_originals: bool):
        self.vault = vault
        self.dry_run = dry_run
        self.keep_originals = keep_originals
        self.own_addresses = own_addresses
        self.index = ConversionIndex(vault / INDEX_FILE)
        self.results: list[ConversionResult] = []

    # -- public
    def convert_file(self, source: Path) -> None:
        try:
            self._convert(source)
        except Exception as error:  # any failure (parse, odd date, path too long) quarantines the file, not the batch
            self.results.append(ConversionResult(source.name, "failed", error=f"{type(error).__name__}: {error}"))
            if not self.dry_run and source.exists():
                self._quarantine(source)

    def _convert(self, source: Path) -> None:
        sha = hashlib.sha256(source.read_bytes()).hexdigest()
        known_note = self.index.note_for_sha(sha)
        if known_note:
            self._file_duplicate(source, known_note)
            return
        parsed = parse_msg_file(source) if source.suffix.lower() == ".msg" else parse_eml_file(source)
        if not parsed.date:
            parsed.date = dt.datetime.fromtimestamp(source.stat().st_mtime).astimezone()
            parsed.date_source = DateSource.FILE_MODIFIED_TIME
        known_note = self.index.note_for_message_id(parsed.message_id)
        if known_note:
            self._file_duplicate(source, known_note)
            return

        note = self._plan_note_path(parsed)
        archive_destination = self._plan_archive(source, note)
        self._write_note(parsed, note, sha, source.suffix.lower().lstrip("."),
                         source_file=self._relative(archive_destination) if archive_destination else "")
        if archive_destination:
            self._move(source, archive_destination)
        self.results.append(ConversionResult(source.name, "converted", note=self._relative(note),
                                             warnings=_all_warnings(parsed)))

    def save_index(self) -> None:
        if self.dry_run:
            return
        self.index.save()
        with (self.vault / LOG_FILE).open("a", encoding="utf-8") as log:
            timestamp = dt.datetime.now().isoformat(timespec="seconds")
            for result in self.results:
                log.write(json.dumps({"at": timestamp, **result.to_dict()}, ensure_ascii=False) + "\n")

    # -- paths
    def _relative(self, path: Path) -> str:
        return path.relative_to(self.vault).as_posix()

    @staticmethod
    def _unique(path: Path) -> Path:
        """`path`, or `name (2).ext`, `name (3).ext`... if it already exists."""
        if not path.exists():
            return path
        for number in range(2, MAX_UNIQUE_SUFFIX):
            candidate = path.with_name(f"{path.stem} ({number}){path.suffix}")
            if not candidate.exists():
                return candidate
        raise RuntimeError(f"cannot find free name for {path}")

    @staticmethod
    def _local_time(parsed: ParsedEmail) -> dt.datetime:
        return parsed.date.astimezone() if parsed.date.tzinfo else parsed.date

    def _plan_note_path(self, parsed: ParsedEmail) -> Path:
        local = self._local_time(parsed)
        sender = parsed.sender[0] if parsed.sender else UNKNOWN_SENDER
        stem = sanitize_filename(f"{local:%Y-%m-%d %H%M} {sender.name} - {normalize_subject(parsed.subject)}")
        return self._unique(self.vault / "raw" / "emails" / f"{local:%Y}" / f"{local:%m}" / f"{stem}.md")

    def _plan_archive(self, source: Path, note: Path) -> Path | None:
        """Where the original will be moved to, or None when originals stay in place."""
        if self.keep_originals or self.dry_run:
            return None
        relative = note.with_suffix("").relative_to(self.vault / "raw" / "emails")
        destination = self.vault / "raw" / "msg" / relative.parent / (relative.name + source.suffix.lower())
        return self._unique(destination)

    # -- writing
    def _write_note(self, parsed: ParsedEmail, note: Path, sha: str, source_format: str, *,
                    source_file: str = "", embedded_in: str = "") -> None:
        local = self._local_time(parsed)
        sender = parsed.sender[0] if parsed.sender else UNKNOWN_SENDER
        attachments_dir = self.vault / "raw" / "attachments" / f"{local:%Y}" / f"{local:%m}" / note.stem
        note_link = self._relative(note.with_suffix(""))

        attachment_lines = self._save_attachments(parsed.attachments, attachments_dir, note_link)
        for embedded in parsed.embedded:  # forwarded-as-attachment emails become their own notes
            if not embedded.date:
                embedded.date, embedded.date_source = parsed.date, DateSource.PARENT_EMAIL
            child = self._plan_note_path(embedded)
            self._write_note(embedded, child, sha, source_format, embedded_in=note_link)
            attachment_lines.append(f"- Embedded email: [[{self._relative(child.with_suffix(''))}]]")

        if self.dry_run:
            return
        direction = "sent" if sender.email in self.own_addresses else "received"
        text = render_note(parsed, sender=sender, local=local, direction=direction, attachment_lines=attachment_lines,
                           source_format=source_format, source_file=source_file, sha=sha, embedded_in=embedded_in)
        note.parent.mkdir(parents=True, exist_ok=True)
        note.write_text(text, encoding="utf-8")
        self.index.add(self._relative(note), message_id=parsed.message_id, sha=None if embedded_in else sha)

    def _save_attachments(self, attachments: list[Attachment], directory: Path, email_link: str) -> list[str]:
        """Write attachment files (+ extracted text); returns one markdown bullet per attachment."""
        lines = []
        for attachment in attachments:
            is_signature_logo = (attachment.inline and Path(attachment.filename).suffix.lower() in IMAGE_EXT
                                 and len(attachment.data) < SMALL_INLINE_IMAGE)
            if is_signature_logo:
                continue
            filename = sanitize_filename(attachment.filename, MAX_ATTACHMENT_NAME)
            blocked = Path(filename).suffix.lower() in BLOCKED_EXT
            too_large = len(attachment.data) > MAX_EXTRACT_BYTES
            if blocked:
                filename += ".blocked"
            elif Path(filename).suffix.lower() in NOTE_LIKE_EXT:
                filename += ".txt"
            target = self._unique(directory / filename)
            line = f"- [[{self._relative(target)}|{filename}]] ({len(attachment.data) / 1024:.0f} KB)"
            if blocked:
                line += " · ⚠ executable type, not opened"
            elif too_large:
                line += " · ⚠ over size limit, text not extracted"
            if not self.dry_run:
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(attachment.data)
                if not (blocked or too_large):
                    line += self._write_extracted_text(target, email_link)
            lines.append(line)
        return lines

    def _write_extracted_text(self, attachment_path: Path, email_link: str) -> str:
        """Writes `<file>.extracted.md` next to the attachment; returns the suffix for its bullet ("" if no text)."""
        text = normalize_text(extract_attachment_text(attachment_path))
        if not text:
            return ""
        clipped = defang_markdown(text[:TEXT_EXTRACT_LIMIT])
        if len(text) > TEXT_EXTRACT_LIMIT:
            clipped += "\n\n… (truncated)"
        extracted = attachment_path.with_name(attachment_path.name + ".extracted.md")
        extracted.write_text(
            "---\ntype: attachment-text\n"
            f"attachment: {yaml_str('[[' + self._relative(attachment_path) + ']]')}\n"
            f"email: {yaml_str('[[' + email_link + ']]')}\n"
            "---\n\n" + clipped + "\n", encoding="utf-8")
        return f" · text: [[{self._relative(extracted.with_suffix(''))}|extracted]]"

    # -- inbox file handling
    def _file_duplicate(self, source: Path, known_note: str) -> None:
        self.results.append(ConversionResult(source.name, "duplicate", note=known_note))
        if self.keep_originals or self.dry_run:
            return
        self._move(source, self._unique(self.vault / "raw" / "msg" / "_duplicates" / source.name))

    def _quarantine(self, source: Path) -> None:
        """Move an unparseable file to 00-inbox/_failed/ with the traceback next to it."""
        failed_dir = self.vault / INBOX_DIRNAME / FAILED_DIRNAME
        failed_dir.mkdir(parents=True, exist_ok=True)
        if source.parent != failed_dir:
            shutil.move(str(source), str(failed_dir / source.name))
        (failed_dir / (source.name + ".error.txt")).write_text(traceback.format_exc(), encoding="utf-8")

    @staticmethod
    def _move(source: Path, destination: Path) -> None:
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(source), str(destination))


def _all_warnings(parsed: ParsedEmail) -> list[str]:
    return parsed.warnings + [warning for embedded in parsed.embedded for warning in _all_warnings(embedded)]


# ---------------------------------------------------------------- command line

def build_parser() -> argparse.ArgumentParser:
    # allow_abbrev=False: settings.json denies the literal `--vault`, so `--vau` must not reach it
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter,
                                     allow_abbrev=False)
    parser.add_argument("files", nargs="*", help="specific .msg/.eml files (default: everything in 00-inbox/)")
    parser.add_argument("--vault", help="vault root (default: $CLAUDE_PROJECT_DIR or nearest folder with CLAUDE.md)")
    parser.add_argument("--me", action="append", default=[], help="your own address(es); also env EMAIL_VAULT_ME")
    parser.add_argument("--keep", action="store_true", help="leave originals in the inbox instead of archiving")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--json", action="store_true", help="print results as JSON")
    return parser


def select_files(inbox: Path, explicit_files: list[str]) -> list[Path]:
    """Files named on the command line (must be inside the inbox), else every .msg/.eml in the inbox."""
    if not explicit_files:
        return sorted(file for file in inbox.rglob("*")
                      if file.suffix.lower() in EMAIL_SUFFIXES and FAILED_DIRNAME not in file.parts)
    inbox_root = inbox.resolve()
    files = [Path(name).resolve() for name in explicit_files]
    outside = [file for file in files if not file.is_relative_to(inbox_root)]
    if outside:
        raise RefusedFilesError("refused: files must be inside 00-inbox/: " + ", ".join(str(file) for file in outside))
    return files


def print_summary(vault: Path, results: list[ConversionResult], setup: list[str], *,
                  dry_run: bool, as_json: bool) -> None:
    counts: dict[str, int] = {}
    for result in results:
        counts[result.status] = counts.get(result.status, 0) + 1
    if as_json:
        payload = {"vault": str(vault), "counts": counts, "setup": setup,
                   "results": [result.to_dict() for result in results]}
        print(json.dumps(payload, ensure_ascii=False, indent=1))
        return
    print(f"Vault: {vault}{'  (dry run)' if dry_run else ''}")
    print("Summary: " + (", ".join(f"{status} {count}" for status, count in counts.items()) or "inbox empty"))
    for warning in setup:
        print(f"Setup warning: {warning}")
    for result in results:
        print(f"  [{result.status}] {result.file}  ->  {result.note or result.error or ''}")


def main(argv: list[str] | None = None) -> int:
    utf8_stdout()
    args = build_parser().parse_args(argv)
    vault = find_vault(args.vault)
    try:
        files = select_files(vault / INBOX_DIRNAME, args.files)
    except RefusedFilesError as refusal:
        print(refusal)
        return 2
    owner = owner_addresses(args.me, vault)
    converter = Converter(vault, owner, dry_run=args.dry_run, keep_originals=args.keep)
    try:
        for file in files:
            converter.convert_file(file)
    finally:  # keep the index of what was converted even if the run is interrupted
        converter.save_index()
    print_summary(vault, converter.results, check_setup(vault, owner), dry_run=args.dry_run, as_json=args.json)
    return 1 if any(result.status == "failed" for result in converter.results) else 0


if __name__ == "__main__":
    sys.exit(main())
